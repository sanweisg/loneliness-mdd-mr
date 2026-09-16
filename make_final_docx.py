#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate final submission DOCX from manuscript_v4.md with embedded figures."""
import sys, os
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX_OUT = r'C:\Users\闫\Desktop\depression_paper_submission\manuscript_final.docx'
FIG_DIR = r'C:\Users\闫\Desktop\depression_paper_submission\R_figs'
MD_IN = r'C:\Users\闫\Desktop\depression_paper_submission\manuscript_v4.md'

fig_map = {
    'Figure 1': 'fig1_forest.png',
    'Figure 2': 'fig2_mediation.png',
    'Figure 3': 'fig3_scatter.png',
    'Figure 4': 'fig4_sensitivity.png',
}

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

def add_heading(text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

def add_para(text, bold=False, italic=False, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parts = __import__('re').split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            r = p.add_run(part)
            r.bold = bold
        r.italic = italic
        r.font.name = 'Times New Roman'
    return p

def add_table_md(rows):
    if not rows:
        return
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = cell
            for para in c.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
                    run.font.name = 'Times New Roman'
                    if i == 0:
                        run.bold = True

with open(MD_IN, encoding='utf-8') as f:
    lines = f.read().split('\n')

i = 0
in_table = False
table_buf = []

while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    # Table detection
    if stripped.startswith('|') and stripped.endswith('|'):
        in_table = True
        table_buf.append(stripped)
        i += 1
        continue
    elif in_table:
        if table_buf:
            rows = []
            for tr in table_buf:
                cells = [c.strip() for c in tr.strip().strip('|').split('|')]
                rows.append(cells)
            rows = [r for r in rows if not all(set(c) <= set('-: ') for c in r)]
            add_table_md(rows)
            doc.add_paragraph()
        table_buf = []
        in_table = False

    if not stripped or stripped == '---':
        i += 1
        continue

    # Heading
    m = __import__('re').match(r'^(#{1,4})\s+(.*)', stripped)
    if m:
        level = min(len(m.group(1)), 4)
        add_heading(m.group(2), level)
        i += 1
        continue

    # Figure legend
    mfig = __import__('re').match(r'^\*\*(Figure \d)\.\*\*\s*(.*)', stripped)
    if mfig:
        add_para(stripped)
        fign = mfig.group(1)
        if fign in fig_map:
            fp = os.path.join(FIG_DIR, fig_map[fign])
            if os.path.exists(fp):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(fp, width=Inches(6.0))
        i += 1
        continue

    # Blockquote / italic
    if stripped.startswith('>'):
        add_para(stripped.lstrip('> ').strip(), italic=True)
        i += 1
        continue

    # Checkbox
    if stripped.startswith('- [ ]') or stripped.startswith('- [x]'):
        add_para(stripped)
        i += 1
        continue

    # List
    if stripped.startswith('- '):
        p = doc.add_paragraph(stripped[2:], style='List Bullet')
        for run in p.runs:
            run.font.name = 'Times New Roman'
        i += 1
        continue
    if __import__('re').match(r'^\d+\.\s', stripped):
        p = doc.add_paragraph(__import__('re').sub(r'^\d+\.\s', '', stripped), style='List Number')
        for run in p.runs:
            run.font.name = 'Times New Roman'
        i += 1
        continue

    add_para(stripped)
    i += 1

if in_table and table_buf:
    rows = []
    for tr in table_buf:
        cells = [c.strip() for c in tr.strip().strip('|').split('|')]
        rows.append(cells)
    rows = [r for r in rows if not all(set(c) <= set('-: ') for c in r)]
    add_table_md(rows)

doc.save(DOCX_OUT)
print(f'Saved: {DOCX_OUT}')
print(f'Size: {os.path.getsize(DOCX_OUT):,} bytes')
